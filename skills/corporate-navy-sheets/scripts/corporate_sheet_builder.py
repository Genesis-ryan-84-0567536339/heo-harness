#!/usr/bin/env python3
"""
Corporate Navy Spreadsheet Builder Helper
Standardized openpyxl styling generator for Antigravity Co-Pilot
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

NAVY_HEADER_FILL = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
WHITE_BOLD_FONT = Font(name="Arial", size=11, bold=True, color="FFFFFF")
ZEBRA_EVEN_FILL = PatternFill(start_color="F0F4F8", end_color="F0F4F8", fill_type="solid")
ZEBRA_ODD_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
TOTAL_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
BOLD_FONT = Font(name="Arial", size=11, bold=True, color="1B365D")
REGULAR_FONT = Font(name="Arial", size=10, color="000000")

THIN_BORDER_SIDE = Side(border_style="thin", color="D0D7DE")
DOUBLE_BOTTOM_SIDE = Side(border_style="double", color="1B365D")
THIN_TOP_SIDE = Side(border_style="thin", color="1B365D")

DATA_BORDER = Border(
    left=THIN_BORDER_SIDE,
    right=THIN_BORDER_SIDE,
    top=THIN_BORDER_SIDE,
    bottom=THIN_BORDER_SIDE
)

TOTAL_BORDER = Border(
    left=THIN_BORDER_SIDE,
    right=THIN_BORDER_SIDE,
    top=THIN_TOP_SIDE,
    bottom=DOUBLE_BOTTOM_SIDE
)

def style_corporate_sheet(ws, title, headers, data, total_labels=None, sum_cols=None):
    """
    ws: worksheet
    title: String sheet title
    headers: list of column header strings
    data: list of rows (each row is a list/tuple of values)
    total_labels: String for total row label, e.g. "Tổng cộng"
    sum_cols: list of column indices (1-based) to apply =SUM()
    """
    # Title row
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = ws.cell(row=1, column=1, value=title.upper())
    title_cell.font = Font(name="Arial", size=14, bold=True, color="1B365D")
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 35

    # Subtitle / spacer
    ws.row_dimensions[2].height = 10

    # Headers
    start_row = 3
    ws.row_dimensions[start_row].height = 28
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=h)
        cell.font = WHITE_BOLD_FONT
        cell.fill = NAVY_HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = DATA_BORDER

    # Data Rows
    current_row = start_row + 1
    for r_i, row_data in enumerate(data):
        ws.row_dimensions[current_row].height = 22
        fill = ZEBRA_EVEN_FILL if r_i % 2 == 0 else ZEBRA_ODD_FILL
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = REGULAR_FONT
            cell.fill = fill
            cell.border = DATA_BORDER

            # Auto-align
            if isinstance(val, (int, float)):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '#,##0'
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
        current_row += 1

    # Total Row (if requested)
    if total_labels and sum_cols:
        ws.row_dimensions[current_row].height = 25
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.font = BOLD_FONT
            cell.fill = TOTAL_FILL
            cell.border = TOTAL_BORDER
            if col_idx == 1:
                cell.value = total_labels
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif col_idx in sum_cols:
                col_letter = get_column_letter(col_idx)
                cell.value = f"=SUM({col_letter}{start_row + 1}:{col_letter}{current_row - 1})"
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '#,##0'
        current_row += 1

    # Auto-fit columns
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or '')
            if cell.row == 1:
                continue
            max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)
